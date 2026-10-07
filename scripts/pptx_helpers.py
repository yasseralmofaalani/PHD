import os
import sys
import json
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# --- COLOR PALETTE ---
COLOR_DEEP_TEAL = RGBColor(66, 129, 119)     # #428177 - Primary Brand Teal
COLOR_DARK_TEAL = RGBColor(20, 55, 50)       # #143732 - Dark Executive Teal
COLOR_BURGUNDY  = RGBColor(107, 31, 42)      # #6B1F2A - Secondary Academic Red
COLOR_DARK_SLATE= RGBColor(15, 23, 42)       # #0F172A - Deep Slate Background
COLOR_SLATE_CARD= RGBColor(30, 41, 59)       # #1E293B - Dark Mode Card
COLOR_LIGHT_BG  = RGBColor(246, 248, 250)    # Clean warm off-white
COLOR_WHITE     = RGBColor(255, 255, 255)
COLOR_TEXT_DARK = RGBColor(15, 23, 42)
COLOR_TEXT_MUTED= RGBColor(100, 116, 139)
COLOR_BORDER    = RGBColor(226, 232, 240)
COLOR_GOLD      = RGBColor(217, 119, 6)       # #D97706
COLOR_EMERALD   = RGBColor(16, 185, 129)     # #10B981
COLOR_BLUE      = RGBColor(37, 99, 235)      # #2563EB

FONT_TITLE = "Segoe UI"
FONT_BODY  = "Segoe UI"

def set_slide_background(slide, prs, color):
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    return bg

def add_header(slide, title_ar, section_name, slide_num, dark=False):
    # Top bar badge
    tb_badge = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
    tf_badge = tb_badge.text_frame
    tf_badge.word_wrap = True
    p_badge = tf_badge.paragraphs[0]
    p_badge.text = f"المعهد العالي للعلوم التطبيقية والتكنولوجيا  |  {section_name}"
    p_badge.font.name = FONT_BODY
    p_badge.font.size = Pt(11)
    p_badge.font.bold = True
    p_badge.font.color.rgb = COLOR_GOLD if dark else COLOR_DEEP_TEAL
    p_badge.alignment = PP_ALIGN.RIGHT

    # Main Title
    tb_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.8))
    tf_title = tb_title.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = title_ar
    p_title.font.name = FONT_TITLE
    p_title.font.size = Pt(22)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_WHITE if dark else COLOR_TEXT_DARK
    p_title.alignment = PP_ALIGN.RIGHT

def add_footer(slide, current_num, total_slides=73, dark=False):
    tb_footer = slide.shapes.add_textbox(Inches(0.8), Inches(6.9), Inches(11.7), Inches(0.4))
    tf = tb_footer.text_frame
    p = tf.paragraphs[0]
    p.text = f"أطروحة دكتوراه: التخطيط والتحكم الذكي بالشبكات الخلوية في بيئة GIS  —  ياسر المفعلاني   |   شريحة {current_num} من {total_slides}"
    p.font.name = FONT_BODY
    p.font.size = Pt(9.5)
    p.font.color.rgb = RGBColor(148, 163, 184) if dark else COLOR_TEXT_MUTED
    p.alignment = PP_ALIGN.RIGHT

def set_speaker_notes(slide, notes_dict):
    if not notes_dict:
        return
    notes_slide = slide.notes_slide
    tf = notes_slide.notes_text_frame
    tf.clear()
    
    title = notes_dict.get("title", "")
    time_min = notes_dict.get("timeMinutes", "")
    points = notes_dict.get("points", [])
    
    p0 = tf.paragraphs[0]
    p0.text = f"=== {title} (الزمن المقترح: {time_min} دقيقة) ==="
    p0.font.bold = True
    
    for pt in points:
        p = tf.add_paragraph()
        p.text = f"• {pt}"

def add_card(slide, left, top, width, height, title, items, bg_color=COLOR_WHITE, border_color=COLOR_BORDER, title_color=COLOR_DEEP_TEAL, is_dark=False):
    # Shape container
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)
    else:
        shape.line.fill.background()

    # Text frame
    tb = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), width - Inches(0.4), height - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True

    if title:
        p_title = tf.paragraphs[0]
        p_title.text = title
        p_title.font.name = FONT_TITLE
        p_title.font.size = Pt(15)
        p_title.font.bold = True
        p_title.font.color.rgb = title_color
        p_title.alignment = PP_ALIGN.RIGHT

    for item in items:
        p = tf.add_paragraph()
        p.text = f"• {item}"
        p.font.name = FONT_BODY
        p.font.size = Pt(12)
        p.font.color.rgb = RGBColor(226, 232, 240) if is_dark else COLOR_TEXT_DARK
        p.alignment = PP_ALIGN.RIGHT
        p.space_after = Pt(4)

def add_kpis(slide, top, kpis, is_dark=False):
    # kpis is list of dict: [{"val": "95.12%", "label": "نسبة التغطية الكلية", "sub": "BPSO تفوق على AGA", "color": COLOR_DEEP_TEAL}]
    n = len(kpis)
    total_w = 11.7
    gap = 0.25
    card_w = (total_w - (n - 1) * gap) / n
    left_start = 0.8
    
    for i, kpi in enumerate(kpis):
        left = Inches(left_start + i * (card_w + gap))
        h = Inches(2.0)
        
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(card_w), h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = COLOR_SLATE_CARD if is_dark else COLOR_WHITE
        shape.line.color.rgb = kpi.get("color", COLOR_DEEP_TEAL)
        shape.line.width = Pt(2)
        
        tb = slide.shapes.add_textbox(left + Inches(0.15), top + Inches(0.15), Inches(card_w - 0.3), h - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True
        
        # Value
        p_val = tf.paragraphs[0]
        p_val.text = kpi.get("val", "")
        p_val.font.name = FONT_TITLE
        p_val.font.size = Pt(24)
        p_val.font.bold = True
        p_val.font.color.rgb = kpi.get("color", COLOR_DEEP_TEAL)
        p_val.alignment = PP_ALIGN.CENTER
        
        # Label
        p_lbl = tf.add_paragraph()
        p_lbl.text = kpi.get("label", "")
        p_lbl.font.name = FONT_BODY
        p_lbl.font.size = Pt(12)
        p_lbl.font.bold = True
        p_lbl.font.color.rgb = COLOR_WHITE if is_dark else COLOR_TEXT_DARK
        p_lbl.alignment = PP_ALIGN.CENTER
        p_lbl.space_after = Pt(2)
        
        # Subtitle
        if kpi.get("sub"):
            p_sub = tf.add_paragraph()
            p_sub.text = kpi.get("sub", "")
            p_sub.font.name = FONT_BODY
            p_sub.font.size = Pt(10)
            p_sub.font.color.rgb = COLOR_GOLD if is_dark else COLOR_TEXT_MUTED
            p_sub.alignment = PP_ALIGN.CENTER

print("Builder helper definitions loaded successfully.")

