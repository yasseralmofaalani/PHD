# -*- coding: utf-8 -*-
"""
Chapter 4 defense deck — numbers only from the thesis chapter.
Output:
  Chapter_4_5G_Optimization_Defense.pptx
  deliverables/ch4_charts/*.png
  deliverables/Chapter_4_Speaker_Notes.txt
"""
from __future__ import annotations

import os
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.ticker import MultipleLocator
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import nsmap
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
CHARTS = ROOT / "deliverables" / "ch4_charts"
NOTES_PATH = ROOT / "deliverables" / "Chapter_4_Speaker_Notes.txt"
PPTX_PATH = ROOT / "Chapter_4_5G_Optimization_Defense.pptx"

# Palette — fixed algorithm colors
BPSO = RGBColor(0x42, 0x81, 0x77)
AGA = RGBColor(0x6B, 0x1F, 0x2A)
STD = RGBColor(0x8A, 0x92, 0x98)
RND = RGBColor(0xB8, 0xBE, 0xC3)
INK = RGBColor(0x0F, 0x17, 0x2A)
MUTED = RGBColor(0x64, 0x74, 0x8B)
PAPER = RGBColor(0xF6, 0xF4, 0xEC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x2E, 0x7D, 0x5B)
GOLD = RGBColor(0xC8, 0x95, 0x1A)
HAIR = RGBColor(0xE2, 0xE0, 0xD4)
RED = RGBColor(0xB0, 0x3A, 0x2E)

HEX = {"bpso": "#428177", "aga": "#6b1f2a", "std": "#8a9298", "rnd": "#b8bec3", "ink": "#0f172a"}

FONT = "Segoe UI"
TOTAL = 16


def set_rtl(tf):
    body = tf._txBody
    bodyPr = body.find(qn("a:bodyPr"))
    if bodyPr is None:
        return
    bodyPr.set("rtlCol", "1")
    for p in tf.paragraphs:
        pPr = p._p.get_or_add_pPr()
        pPr.set("rtl", "1")
        p.alignment = PP_ALIGN.RIGHT


def fill_shape(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def card(slide, l, t, w, h, fill=WHITE, line=HAIR, line_w=1.25):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = line
    sh.line.width = Pt(line_w)
    sh.adjustments[0] = 0.08
    return sh


def tb(slide, l, t, w, h, text, size=16, bold=False, color=INK, align="right", font=FONT):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = font
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = {"right": PP_ALIGN.RIGHT, "center": PP_ALIGN.CENTER, "left": PP_ALIGN.LEFT}[align]
    if align == "right":
        set_rtl(tf)
    return tf


def notes(slide, title, say, msg, avoid, q, a):
    ns = slide.notes_slide.notes_text_frame
    ns.clear()
    ns.paragraphs[0].text = f"=== {title} ==="
    body = (
        f"\n1) ماذا أقول شفهياً؟\n{say}\n"
        f"\n2) الرسالة الأساسية:\n{msg}\n"
        f"\n3) ما الذي يجب ألا أتوسع فيه؟\n{avoid}\n"
        f"\n4) سؤال متوقع:\n{q}\n"
        f"\n5) جواب مختصر:\n{a}\n"
    )
    p = ns.add_paragraph()
    p.text = body


def header(slide, prs, kicker, title, num):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.12), prs.slide_height)
    fill_shape(bar, BPSO)
    tb(slide, Inches(0.4), Inches(0.18), Inches(10.6), Inches(0.28), f"المساهمة الأولى  ·  الفصل الرابع  ·  {kicker}", 11, True, BPSO)
    tb(slide, Inches(0.4), Inches(0.42), Inches(12.2), Inches(0.55), title, 24, True, INK)
    tb(slide, Inches(11.7), Inches(7.12), Inches(1.3), Inches(0.24), f"{num:02d} / {TOTAL:02d}", 10, True, MUTED, "left")
    src = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.38), prs.slide_width, Inches(0.12))
    fill_shape(src, BPSO)


def blank(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    fill_shape(bg, PAPER)
    return slide


def style_charts():
    plt.rcParams.update(
        {
            "font.family": "Segoe UI",
            "axes.spines.top": False,
            "axes.spines.right": False,
            "axes.edgecolor": "#c8c5b8",
            "axes.linewidth": 0.8,
            "axes.grid": False,
            "figure.facecolor": "white",
            "axes.facecolor": "white",
            "savefig.bbox": "tight",
            "savefig.dpi": 200,
        }
    )


def make_charts():
    CHARTS.mkdir(parents=True, exist_ok=True)
    style_charts()

    # 1. Baseline grouped bars
    fig, axes = plt.subplots(1, 3, figsize=(11.2, 3.6))
    names = ["BPSO", "AGA", "Std. GA", "Random"]
    colors = [HEX["bpso"], HEX["aga"], HEX["std"], HEX["rnd"]]
    series = [
        ("Coverage (%)", [95.14, 94.90, 94.75, 88.35], 86, 97),
        ("Cost (M$)", [130.5, 139.6, 143.8, 150.3], 120, 155),
        ("Energy (MWh)", [76.4, 80.9, 84.1, 91.5], 70, 95),
    ]
    for ax, (title, vals, ymin, ymax) in zip(axes, series):
        ax.bar(names, vals, color=colors, width=0.62)
        ax.set_title(title, loc="left", fontsize=11, fontweight="bold", color=HEX["ink"])
        ax.set_ylim(ymin, ymax)
        ax.tick_params(labelsize=8)
        for i, v in enumerate(vals):
            ax.text(i, v + (ymax - ymin) * 0.015, f"{v}", ha="center", va="bottom", fontsize=8, fontweight="bold")
    fig.tight_layout()
    fig.savefig(CHARTS / "01_baselines.png")
    plt.close()

    # 2. Budget vs coverage — real table 12
    fig, ax = plt.subplots(figsize=(8.4, 4.2))
    budgets = [80, 100, 120, 140]
    aga = [88.7, 91.9, 93.8, 94.8]
    bpso = [89.4, 92.8, 94.5, 95.1]
    ax.plot(budgets, aga, "o-", color=HEX["aga"], lw=2.4, ms=7, label="AGA")
    ax.plot(budgets, bpso, "o-", color=HEX["bpso"], lw=2.4, ms=7, label="BPSO")
    ax.axvline(140, color="#c8951a", ls="--", lw=1.2, alpha=0.85)
    ax.text(140.4, 89.2, "Diminishing returns ≈ 140 M$", color="#8a6410", fontsize=9, fontweight="bold")
    ax.set_title("Budget vs Coverage (Table 12)", loc="left", fontsize=12, fontweight="bold")
    ax.set_xlabel("Budget (M$)")
    ax.set_ylabel("Coverage (%)")
    ax.set_xticks(budgets)
    ax.set_ylim(87, 96)
    ax.legend(frameon=False)
    fig.tight_layout()
    fig.savefig(CHARTS / "02_budget_coverage.png")
    plt.close()

    # 3. Stability error bars — real tables 9–10
    fig, ax = plt.subplots(figsize=(8.4, 4.0))
    labels = ["Coverage (%)", "Cost (M$)", "Energy (MWh)"]
    aga_m = [94.90, 139.6, 80.9]
    aga_s = [0.36, 2.5, 1.8]
    bp_m = [95.14, 130.5, 76.4]
    bp_s = [0.42, 2.9, 2.1]
    # normalize each metric to 0-1 of its pair for visual comparison of variability only? No —
    # three panels is clearer.
    plt.close()
    fig, axes = plt.subplots(1, 3, figsize=(10.6, 3.6))
    for ax, lab, am, asd, bm, bsd in zip(axes, labels, aga_m, aga_s, bp_m, bp_s):
        ax.bar([0], [am], yerr=[asd], color=HEX["aga"], width=0.45, capsize=6, label="AGA")
        ax.bar([1], [bm], yerr=[bsd], color=HEX["bpso"], width=0.45, capsize=6, label="BPSO")
        ax.set_xticks([0, 1], ["AGA", "BPSO"])
        ax.set_title(lab, loc="left", fontsize=11, fontweight="bold")
        pad = max(asd, bsd) * 4 + (max(am, bm) - min(am, bm)) * 0.35 + 0.8
        ax.set_ylim(min(am, bm) - pad, max(am, bm) + pad)
        ax.text(0, am, f"{am:.2f}±{asd}", ha="center", va="bottom", fontsize=8)
        ax.text(1, bm, f"{bm:.2f}±{bsd}", ha="center", va="bottom", fontsize=8)
    fig.tight_layout()
    fig.savefig(CHARTS / "03_stability.png")
    plt.close()

    # 4. Conceptual convergence — labeled illustration
    fig, ax = plt.subplots(figsize=(8.4, 4.0))
    xs = list(range(0, 301, 5))
    # Qualitative shapes only — no claimed numeric fitness
    bp = [1 - 0.72 * (2.718 ** (-x / 38.0)) for x in xs]
    ag = [1 - 0.72 * (2.718 ** (-x / 70.0)) for x in xs]
    ax.plot(xs, bp, color=HEX["bpso"], lw=2.6, label="BPSO")
    ax.plot(xs, ag, color=HEX["aga"], lw=2.6, label="AGA")
    ax.axvline(80, color="#94a3b8", ls=":", lw=1)
    ax.axvline(200, color="#c8951a", ls="--", lw=1.2)
    ax.text(82, 0.22, "faster first 80", color=HEX["bpso"], fontsize=8)
    ax.text(204, 0.22, "stable ~200", color="#8a6410", fontsize=8)
    ax.set_xlim(0, 300)
    ax.set_ylim(0.15, 1.05)
    ax.set_yticks([])
    ax.set_xlabel("Iteration")
    ax.set_ylabel("Objective / Fitness (schematic)")
    ax.set_title("Convergence — conceptual illustration of Figure 14 analysis", loc="left", fontsize=11, fontweight="bold")
    ax.legend(frameon=False)
    fig.text(0.01, 0.01, "Not digitized from the original curve. Qualitative only.", fontsize=8, color="#64748b")
    fig.tight_layout()
    fig.savefig(CHARTS / "04_convergence_illustration.png")
    plt.close()

    # 5. Conceptual Pareto
    fig, ax = plt.subplots(figsize=(7.6, 4.4))
    ax.scatter([150.3, 148, 146], [88.35, 87.8, 86.9], c=HEX["rnd"], s=46, label="Random", zorder=2)
    ax.scatter([143.8, 141, 145], [94.75, 93.9, 94.2], c=HEX["std"], s=50, label="Standard GA", zorder=3)
    ax.scatter([139.6, 136, 142, 133, 145], [94.90, 93.6, 94.4, 92.8, 94.1], c=HEX["aga"], s=58, label="AGA", zorder=4)
    ax.scatter([130.5, 128, 134, 126], [95.14, 94.4, 94.8, 93.7], c=HEX["bpso"], s=70, label="BPSO", zorder=5)
    ax.set_xlabel("Cost (schematic axis)")
    ax.set_ylabel("Coverage (schematic axis)")
    ax.set_title("Pareto region — conceptual illustration of Figure 15 analysis", loc="left", fontsize=11, fontweight="bold")
    ax.legend(frameon=False, loc="lower left")
    fig.text(0.01, 0.01, "Not original Pareto coordinates. Layout follows chapter analysis only.", fontsize=8, color="#64748b")
    fig.tight_layout()
    fig.savefig(CHARTS / "05_pareto_illustration.png")
    plt.close()


def add_picture(slide, path, l, t, w, h):
    slide.shapes.add_picture(str(path), l, t, w, h)


def build():
    make_charts()
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1 Challenge
    s = blank(prs)
    header(s, prs, "The Challenge", "تحسين تخطيط نشر الجيل الخامس في البيئات محدودة الموارد", 1)
    card(s, Inches(0.4), Inches(1.15), Inches(6.3), Inches(5.7), INK)
    tb(s, Inches(0.7), Inches(1.7), Inches(5.7), Inches(1.6), "79,268", 72, True, WHITE, "center")
    tb(s, Inches(0.7), Inches(3.35), Inches(5.7), Inches(0.4), "موقعاً خلوياً تشغيلياً", 20, True, RGBColor(0x9E, 0xE0, 0xD6), "center")
    tb(s, Inches(0.7), Inches(4.0), Inches(5.7), Inches(0.7), "تمثيل تخطيطي GIS  ·  ليس رسماً إحداثياً أصلياً", 12, False, RGBColor(0xCB, 0xD5, 0xE1), "center")
    tb(s, Inches(0.7), Inches(5.5), Inches(5.7), Inches(0.9), "أي المواقع نرقّي إلى الجيل الخامس؟", 22, True, WHITE, "center")
    for i, (t, c) in enumerate([("Coverage ↑", GREEN), ("CAPEX ↓", AGA), ("Energy ↓", GOLD)]):
        card(s, Inches(7.0), Inches(1.15 + i * 1.35), Inches(5.8), Inches(1.2), WHITE, c, 2.25)
        tb(s, Inches(7.25), Inches(1.35 + i * 1.35), Inches(5.3), Inches(0.8), t, 32, True, c, "right")
    notes(
        s,
        "التحدي",
        "لدينا شبكة تشغيلية حقيقية ضخمة وموارد محدودة. السؤال: أي المواقع نرقّي إلى 5G لتعظيم التغطية بأقل كلفة وطاقة.",
        "المسألة هندسية وواقعية وليست تجريبية اصطناعية.",
        "لا نتائج ولا خوارزميات في هذه الشريحة. الخريطة تخطيطية.",
        "هل البيانات حقيقية؟",
        "نعم. MTN وSyriatel بعد تنظيف ومعالجة GIS. الرقم من الجدول 3.",
    )

    # 2 Pipeline
    s = blank(prs)
    header(s, prs, "Research path", "من الشبكة الحقيقية إلى قرار التحسين", 2)
    steps = ["Real Data", "Cleaning", "GIS", "Model", "AGA / BPSO", "Experiments", "Decision"]
    ar = ["بيانات تشغيلية", "تنظيف", "معالجة GIS", "نموذج رياضي", "خوارزميتان", "تجارب", "قرار هندسي"]
    for i, (e, a) in enumerate(zip(steps, ar)):
        x = 0.35 + i * 1.85
        card(s, Inches(x), Inches(2.3), Inches(1.7), Inches(2.0), BPSO if i == 6 else WHITE, BPSO, 1.75)
        tb(s, Inches(x + 0.08), Inches(2.45), Inches(1.54), Inches(0.4), f"{i+1:02d}", 16, True, WHITE if i == 6 else BPSO, "center")
        tb(s, Inches(x + 0.08), Inches(2.9), Inches(1.54), Inches(0.7), a, 14, True, WHITE if i == 6 else INK, "center")
        tb(s, Inches(x + 0.08), Inches(3.6), Inches(1.54), Inches(0.45), e, 11, False, WHITE if i == 6 else MUTED, "center")
    notes(s, "المسار", "اعرض السلسلة كمسار قرار واحد.", "البحث ليس خوارزمية معلّقة.", "لا معادلات هنا.", "أين تبدأ المساهمة؟", "من تجهيز البيانات الحقيقية وصياغة المسألة.")

    # 3 Dataset
    s = blank(prs)
    header(s, prs, "Table 3", "قاعدة البيانات التشغيلية للشبكة السورية", 3)
    kpis = [
        ("79,268", "Total Sites", BPSO),
        ("21,356", "2G", RGBColor(0x7E, 0x57, 0xC2)),
        ("27,902", "3G", RGBColor(0x43, 0xA0, 0x47)),
        ("30,010", "4G Candidates", RGBColor(0x1E, 0x88, 0xE5)),
    ]
    for i, (v, l, c) in enumerate(kpis):
        card(s, Inches(0.4 + i * 3.2), Inches(1.2), Inches(3.05), Inches(1.7), WHITE, c, 2)
        tb(s, Inches(0.5 + i * 3.2), Inches(1.35), Inches(2.85), Inches(0.85), v, 32, True, c, "center")
        tb(s, Inches(0.5 + i * 3.2), Inches(2.2), Inches(2.85), Inches(0.45), l, 14, True, INK, "center")
    extras = [("2 Operators", "MTN · Syriatel"), ("62% Urban", "حضري"), ("38% Rural", "ريفي")]
    for i, (v, l) in enumerate(extras):
        card(s, Inches(0.4 + i * 4.25), Inches(3.15), Inches(4.05), Inches(1.35))
        tb(s, Inches(0.55 + i * 4.25), Inches(3.3), Inches(3.75), Inches(0.55), v, 22, True, BPSO, "center")
        tb(s, Inches(0.55 + i * 4.25), Inches(3.9), Inches(3.75), Inches(0.4), l, 14, False, MUTED, "center")
    card(s, Inches(0.4), Inches(4.7), Inches(12.5), Inches(1.55), RGBColor(0xE8, 0xF3, 0xFC), RGBColor(0x1E, 0x88, 0xE5), 1.75)
    tb(s, Inches(0.65), Inches(4.9), Inches(12.1), Inches(0.45), "4G هي قاعدة مرشحي الترقية في المعمارية غير المستقلة NSA", 20, True, RGBColor(0x1E, 0x88, 0xE5))
    tb(s, Inches(0.65), Inches(5.4), Inches(12.1), Inches(0.55), "N = 30,010 موقع LTE  ·  المصدر: الجدول 3", 14, False, MUTED)
    notes(s, "البيانات", "اقرأ أرقام الجدول 3 كما هي.", "4G = مجموعة القرار.", "لا تدّعِ أن الخريطة إحداثيات أصلية.", "لماذا 4G؟", "الأكثر جاهزية لـ 5G وفق NSA و3GPP.")

    # 4 Why hard
    s = blank(prs)
    header(s, prs, "Problem hardness", "لماذا لا يكفي اختيار مواقع؟", 4)
    hard = [
        ("حجم هائل", "79,268 موقعاً · 30,010 مرشحاً"),
        ("قرار ثنائي", "xᵢ ∈ {0,1}"),
        ("أهداف متعددة", "تغطية × كلفة × طاقة"),
        ("قيود صلبة", "B · K · Cmin"),
        ("غير خطّي", "Min-Max ثم دالة موزونة"),
        ("NP-hard", "الحل الدقيق غير عملي"),
    ]
    for i, (t, d) in enumerate(hard):
        r, c = divmod(i, 3)
        card(s, Inches(0.4 + c * 4.25), Inches(1.2 + r * 1.55), Inches(4.05), Inches(1.4))
        tb(s, Inches(0.55 + c * 4.25), Inches(1.32 + r * 1.55), Inches(3.75), Inches(0.45), t, 18, True, BPSO)
        tb(s, Inches(0.55 + c * 4.25), Inches(1.8 + r * 1.55), Inches(3.75), Inches(0.5), d, 14, False, MUTED)
    card(s, Inches(0.4), Inches(4.45), Inches(12.5), Inches(1.8), INK)
    tb(
        s,
        Inches(0.65),
        Inches(4.65),
        Inches(12.1),
        Inches(1.4),
        "نحن لا نختار مواقع فحسب — بل نبحث في فضاء توافقي هائل (2^30,010) تحت أهداف متنافسة وقيود صلبة. لذلك الاستدلال الفوقي ضرورة.",
        18,
        True,
        WHITE,
    )
    notes(s, "الصعوبة", "أكد أن الفضاء توافقي وأن MILP غير عملي.", "Metaheuristics ضرورة.", "لا إثبات NP-hard.", "لماذا لا Exact؟", "الفصل: الحل الدقيق غير عملي على هذا النطاق.")

    # 5 Model
    s = blank(prs)
    header(s, prs, "Mathematical model", "النموذج الرياضي: قرار واحد وثلاثة أهداف", 5)
    card(s, Inches(0.4), Inches(1.2), Inches(12.5), Inches(1.5), INK)
    tb(s, Inches(0.6), Inches(1.4), Inches(12.1), Inches(0.55), "max  F  =  w1 Coverage  −  w2 Cost  −  w3 Energy", 24, True, WHITE, "center")
    tb(s, Inches(0.6), Inches(2.0), Inches(12.1), Inches(0.4), "w1 + w2 + w3 = 1     ·     Min-Max Normalization → {0,1}", 16, False, RGBColor(0xCB, 0xD5, 0xE1), "center")
    for i, (t, d, c) in enumerate(
        [("Maximize Coverage", "تعظيم التغطية", GREEN), ("Minimize Cost", "تقليل الكلفة", AGA), ("Minimize Energy", "ترشيد الطاقة", GOLD)]
    ):
        card(s, Inches(0.4 + i * 4.25), Inches(2.95), Inches(4.05), Inches(1.45), WHITE, c, 2)
        tb(s, Inches(0.55 + i * 4.25), Inches(3.1), Inches(3.75), Inches(0.45), t, 16, True, c, "center")
        tb(s, Inches(0.55 + i * 4.25), Inches(3.6), Inches(3.75), Inches(0.45), d, 16, True, INK, "center")
    for i, (t, f) in enumerate([("Budget", "Σ Costᵢ xᵢ  ≤  B"), ("Maximum sites K", "Σ xᵢ  ≤  K"), ("Minimum coverage", "Coverage(x)  ≥  Cmin")]):
        card(s, Inches(0.4 + i * 4.25), Inches(4.6), Inches(4.05), Inches(1.55))
        tb(s, Inches(0.55 + i * 4.25), Inches(4.75), Inches(3.75), Inches(0.4), t, 14, True, MUTED, "center")
        tb(s, Inches(0.55 + i * 4.25), Inches(5.2), Inches(3.75), Inches(0.6), f, 16, True, INK, "center")
    notes(s, "النموذج", "قرار ثنائي + ثلاثة أهداف + ثلاثة قيود.", "الأوزان مجموعها واحد بعد التوحيد.", "لا تشتق الأوزان.", "كيف دُمجت الوحدات؟", "Min-Max إلى قيم لا بُعدية.")

    # 6 Two algorithms
    s = blank(prs)
    header(s, prs, "AGA vs BPSO", "لماذا خوارزميتان؟ تطور مقابل ذكاء سربي", 6)
    card(s, Inches(0.4), Inches(1.2), Inches(6.15), Inches(5.0), RGBColor(0xF8, 0xEE, 0xF0), AGA, 2)
    tb(s, Inches(0.6), Inches(1.4), Inches(5.75), Inches(0.5), "AGA  ·  Evolution", 26, True, AGA, "center")
    for i, t in enumerate(["Population-based", "Tournament selection k=3", "Two-point crossover 0.85", "Adaptive mutation 0.08 → 0.02", "Constraint repair"]):
        tb(s, Inches(0.8), Inches(2.15 + i * 0.7), Inches(5.4), Inches(0.55), t, 16, True, INK)
    card(s, Inches(6.8), Inches(1.2), Inches(6.15), Inches(5.0), RGBColor(0xE8, 0xF3, 0xF1), BPSO, 2)
    tb(s, Inches(7.0), Inches(1.4), Inches(5.75), Inches(0.5), "BPSO  ·  Swarm", 26, True, BPSO, "center")
    for i, t in enumerate(["Swarm-based", "pBest + gBest", "Velocity update", "Sigmoid → binary position", "Constraint repair"]):
        tb(s, Inches(7.2), Inches(2.15 + i * 0.7), Inches(5.4), Inches(0.55), t, 16, True, INK)
    notes(s, "الخوارزميتان", "استراتيجيتان مستقلتان على البيانات والقيود نفسها.", "لا فائز هنا.", "لا تقرأ كل المعاملات.", "لماذا هاتان؟", "مدرستان مركزيتان تناسبان القرار الثنائي واسع النطاق.")

    # 7 Repair
    s = blank(prs)
    header(s, prs, "Key innovation", "كيف نمنع الخوارزمية من إنتاج حلول غير صالحة؟", 7)
    stages = [
        ("Infeasible", "تجاوز الميزانية أو تغطية دون الحد", RED),
        ("Adaptive Repair", "حذف الأدنى جدوى أو إضافة الأعلى جدوى", GOLD),
        ("Feasible", "Budget ✓   Coverage ✓   Valid ✓", GREEN),
    ]
    for i, (t, d, c) in enumerate(stages):
        card(s, Inches(0.4 + i * 4.25), Inches(1.2), Inches(4.05), Inches(1.7), WHITE, c, 2.25)
        tb(s, Inches(0.55 + i * 4.25), Inches(1.35), Inches(3.75), Inches(0.5), t, 20, True, c, "center")
        tb(s, Inches(0.55 + i * 4.25), Inches(1.95), Inches(3.75), Inches(0.7), d, 14, True, INK, "center")
    card(s, Inches(0.4), Inches(3.15), Inches(6.15), Inches(3.05), RGBColor(0xFD, 0xF2, 0xF2), RED, 1.5)
    tb(s, Inches(0.6), Inches(3.35), Inches(5.75), Inches(0.45), "إذا تجاوزت الميزانية B", 18, True, RED)
    tb(s, Inches(0.6), Inches(3.9), Inches(5.75), Inches(1.8), "إزالة المواقع ذات أدنى مساهمة (تغطية / تكلفة) حتى الالتزام بالقيد.", 16, False, INK)
    card(s, Inches(6.8), Inches(3.15), Inches(6.15), Inches(3.05), RGBColor(0xFD, 0xF7, 0xE8), GOLD, 1.5)
    tb(s, Inches(7.0), Inches(3.35), Inches(5.75), Inches(0.45), "إذا كانت التغطية دون Cmin", 18, True, GOLD)
    tb(s, Inches(7.0), Inches(3.9), Inches(5.75), Inches(1.8), "إضافة المواقع ذات أعلى مساهمة (تغطية / تكلفة) حتى بلوغ الحد الأدنى.", 16, False, INK)
    notes(s, "الإصلاح", "الحلول غير الصالحة تُصلَح ولا تُحذف. الآلية مشتركة.", "هذا جزء من المنهجية.", "لا تفتح قيد العدالة.", "ما أهميته؟", "يضمن حلولاً مقبولة بعد كل تحديث.")

    # 8 Experiments
    s = blank(prs)
    header(s, prs, "Experimental setup", "إعداد التجارب: مقارنة عادلة وقابلة للتكرار", 8)
    dash = [("Sites", "79,268"), ("Operators", "2"), ("Algorithms", "AGA + BPSO"), ("Pop / Swarm", "150"), ("Iterations", "300"), ("Runs", "15")]
    for i, (l, v) in enumerate(dash):
        card(s, Inches(0.35 + i * 2.15), Inches(1.2), Inches(2.05), Inches(1.45))
        tb(s, Inches(0.4 + i * 2.15), Inches(1.3), Inches(1.95), Inches(0.35), l, 11, True, MUTED, "center")
        tb(s, Inches(0.4 + i * 2.15), Inches(1.7), Inches(1.95), Inches(0.7), v, 18, True, BPSO, "center")
    for i, (sid, sites, bud) in enumerate([("S1", "5,000", "منخفضة"), ("S2", "10,000", "متوسطة"), ("S3", "15,000", "مرتفعة")]):
        card(s, Inches(0.4 + i * 4.25), Inches(2.9), Inches(4.05), Inches(2.0), WHITE, BPSO if i == 1 else GOLD if i == 0 else INK, 2)
        tb(s, Inches(0.55 + i * 4.25), Inches(3.05), Inches(3.75), Inches(0.4), sid, 20, True, BPSO if i == 1 else INK, "center")
        tb(s, Inches(0.55 + i * 4.25), Inches(3.5), Inches(3.75), Inches(0.5), sites + " sites", 22, True, INK, "center")
        tb(s, Inches(0.55 + i * 4.25), Inches(4.1), Inches(3.75), Inches(0.4), "ميزانية " + bud, 14, False, MUTED, "center")
    tb(s, Inches(0.4), Inches(5.15), Inches(12.5), Inches(0.9), "Baselines: Proposed AGA  ·  Proposed BPSO  ·  Standard GA (w/o Repair)  ·  Random Selection\nالنتائج التفصيلية المعروضة للسيناريو S2", 14, False, MUTED)
    notes(s, "التجارب", "150 / 300 / 15 وسيناريوهات 5k–15k.", "تكافؤ حسابي.", "لا تخلط S4/S5.", "لماذا 15 run؟", "متوسط وانحراف عبر بذور مختلفة.")

    # 9 Hero
    s = blank(prs)
    header(s, prs, "Table 11 · S2", "BPSO Achieves the Best Overall Performance", 9)
    card(s, Inches(0.4), Inches(1.2), Inches(6.15), Inches(3.55), RGBColor(0xE8, 0xF3, 0xF1), BPSO, 2.5)
    tb(s, Inches(0.6), Inches(1.35), Inches(5.75), Inches(0.4), "BPSO", 20, True, BPSO, "center")
    tb(s, Inches(0.6), Inches(1.75), Inches(5.75), Inches(1.0), "95.14%", 54, True, BPSO, "center")
    tb(s, Inches(0.6), Inches(2.85), Inches(5.75), Inches(1.4), "130.5 M$     ·     76.4 MWh", 20, True, INK, "center")
    card(s, Inches(6.8), Inches(1.2), Inches(6.15), Inches(3.55), RGBColor(0xF8, 0xEE, 0xF0), AGA, 2.5)
    tb(s, Inches(7.0), Inches(1.35), Inches(5.75), Inches(0.4), "AGA", 20, True, AGA, "center")
    tb(s, Inches(7.0), Inches(1.75), Inches(5.75), Inches(1.0), "94.90%", 54, True, AGA, "center")
    tb(s, Inches(7.0), Inches(2.85), Inches(5.75), Inches(1.4), "139.6 M$     ·     80.9 MWh", 20, True, INK, "center")
    for i, (t, v) in enumerate([("Coverage", "+0.24%"), ("Lower Cost", "~6.5%"), ("Better Energy", "~5.5%")]):
        card(s, Inches(0.4 + i * 4.25), Inches(4.95), Inches(4.05), Inches(1.25), WHITE, GREEN, 2)
        tb(s, Inches(0.55 + i * 4.25), Inches(5.05), Inches(3.75), Inches(0.35), t, 13, True, MUTED, "center")
        tb(s, Inches(0.55 + i * 4.25), Inches(5.4), Inches(3.75), Inches(0.55), v, 26, True, GREEN, "center")
    notes(s, "النتيجة", "اقرأ الأرقام ببطء ثم الفروق الثلاثة.", "BPSO أفضل إجمالاً في S2.", "لا تُلغِ استقرار AGA.", "هل 0.24% مهم؟", "الفصل يعرضه ارتفاعاً طفيفاً مع وفر أوضح في الكلفة والطاقة.")

    # 10 Stability
    s = blank(prs)
    header(s, prs, "Tables 9–10", "الأداء وحده لا يكفي: تحليل الاستقرار", 10)
    add_picture(s, CHARTS / "03_stability.png", Inches(0.35), Inches(1.15), Inches(8.3), Inches(4.7))
    card(s, Inches(8.8), Inches(1.2), Inches(4.1), Inches(2.15), RGBColor(0xE8, 0xF3, 0xF1), BPSO, 2)
    tb(s, Inches(8.95), Inches(1.4), Inches(3.8), Inches(1.7), "BPSO\nأفضل متوسط أداء", 20, True, BPSO)
    card(s, Inches(8.8), Inches(3.55), Inches(4.1), Inches(2.15), RGBColor(0xF8, 0xEE, 0xF0), AGA, 2)
    tb(s, Inches(8.95), Inches(3.75), Inches(3.8), Inches(1.7), "AGA\nاستقرار أعلى / تشتت أقل", 20, True, AGA)
    notes(s, "الاستقرار", "15 تشغيلاً: BPSO أفضل وسطياً وAGA أقل انحرافاً.", "الأداء ≠ الثبات.", "لا اختبار فرضيات غير مذكور.", "لماذا AGA أثبت؟", "تطور متدرج وطفرة تكيفية.")

    # 11 Convergence
    s = blank(prs)
    header(s, prs, "التقارب: سرعة الوصول مقابل سلاسة التحسن", 11)
    add_picture(s, CHARTS / "04_convergence_illustration.png", Inches(0.35), Inches(1.15), Inches(8.5), Inches(5.0))
    card(s, Inches(9.0), Inches(1.2), Inches(3.9), Inches(5.0))
    for i, t in enumerate(["BPSO أسرع في أول 80 تكراراً", "استقرار بعد نحو 200", "T = 300 كافٍ عملياً", "AGA أكثر سلاسة"]):
        tb(s, Inches(9.15), Inches(1.45 + i * 1.05), Inches(3.6), Inches(0.9), t, 16, True, BPSO if i == 0 else INK)
    notes(s, "التقارب", "الرسم توضيحي وفق تحليل الشكل 14.", "لا نقاط مرقمنة.", "لا تخترع لياقة رقمية.", "لماذا 300؟", "الاستقرار بعد ~200 يجعل 300 كافياً.")

    # 12 Pareto
    s = blank(prs)
    header(s, prs, "Figure 15 analysis", "جبهة باريتو: التغطية مقابل التكلفة", 12)
    add_picture(s, CHARTS / "05_pareto_illustration.png", Inches(0.3), Inches(1.15), Inches(8.2), Inches(5.1))
    card(s, Inches(8.7), Inches(1.2), Inches(4.2), Inches(5.0), INK)
    tb(
        s,
        Inches(8.9),
        Inches(1.6),
        Inches(3.85),
        Inches(4.2),
        "حلول BPSO أقرب عموماً إلى منطقة المفاضلة المثلى.\n\nAGA تمنح تنوعاً أوسع على الجبهة.\n\nكلاهما يتفوق على GA التقليدية والعشوائي.",
        16,
        True,
        WHITE,
    )
    notes(s, "باريتو", "توضيح مفهومي للشكل 15.", "ليست إحداثيات أصلية.", "لا تدّعِ نقاطاً مقيسة.", "ما معنى الجبهة؟", "حلول غير مغلوبة بين التغطية والكلفة.")

    # 13 Budget
    s = blank(prs)
    header(s, prs, "Table 12 · Figure 16", "الميزانية مقابل التغطية: عائد متناقص", 13)
    add_picture(s, CHARTS / "02_budget_coverage.png", Inches(0.3), Inches(1.15), Inches(8.5), Inches(5.1))
    card(s, Inches(8.95), Inches(1.2), Inches(3.95), Inches(5.0), GOLD)
    tb(s, Inches(9.1), Inches(1.5), Inches(3.65), Inches(0.5), "Diminishing Returns", 16, True, WHITE, "center")
    tb(s, Inches(9.1), Inches(2.15), Inches(3.65), Inches(0.8), "بعد ≈ 140 M$", 28, True, WHITE, "center")
    tb(s, Inches(9.1), Inches(3.2), Inches(3.65), Inches(2.4), "زيادة الميزانية بعد هذا الحد لا تنتج تحسناً كبيراً في التغطية. حد جدوى هندسي.", 16, True, WHITE)
    notes(s, "الميزانية", "BPSO أعلى تغطية عند كل مستوى. بعد 140 العائد يتناقص.", "قرار تمويلي.", "لا تعمم الرقم.", "هل 140 موصى بها؟", "الفصل: ما بعدها لا يشتري تغطية كبيرة.")

    # 14 Baselines
    s = blank(prs)
    header(s, prs, "Tables 11 & 13", "المقارنة مع خطوط الأساس", 14)
    add_picture(s, CHARTS / "01_baselines.png", Inches(0.35), Inches(1.15), Inches(12.6), Inches(4.35))
    tb(s, Inches(0.4), Inches(5.6), Inches(12.5), Inches(0.7), "Runtime: BPSO 118 s   ·   AGA 142 s   ·   Std. GA 118 s   ·   Random 25 s", 16, True, MUTED, "center")
    notes(s, "الأساس", "العشوائي ضعيف بوضوح. المقترحتان أعلى جودة.", "الإصلاح له أثر.", "زمن Random لا يعني جودة.", "لماذا Std.GA = 118؟", "كما في الجدول 13.")

    # 15 Decision
    s = blank(prs)
    header(s, prs, "Table 14", "أي خوارزمية نختار؟", 15)
    rows = [
        ("Criterion", "AGA", "BPSO"),
        ("Coverage", "Good", "Best"),
        ("Cost", "Higher", "Lower"),
        ("Energy", "Higher", "Lower"),
        ("Convergence", "Slower", "Faster"),
        ("Stability", "Better", "Lower"),
        ("Runtime", "142 s", "118 s"),
        ("Overall", "Strong", "Preferred"),
    ]
    for i, (a, b, c) in enumerate(rows):
        y = 1.2 + i * 0.55
        fill = INK if i == 0 else (WHITE if i % 2 else RGBColor(0xF3, 0xF1, 0xE8))
        card(s, Inches(0.4), Inches(y), Inches(7.5), Inches(0.5), fill, HAIR, 0.75)
        col = WHITE if i == 0 else INK
        tb(s, Inches(0.5), Inches(y + 0.05), Inches(2.3), Inches(0.4), a, 14, True, col)
        tb(s, Inches(2.9), Inches(y + 0.05), Inches(2.3), Inches(0.4), b, 14, True, WHITE if i == 0 else AGA, "center")
        tb(s, Inches(5.3), Inches(y + 0.05), Inches(2.4), Inches(0.4), c, 14, True, WHITE if i == 0 else BPSO, "center")
    card(s, Inches(8.15), Inches(1.2), Inches(4.75), Inches(2.7), BPSO)
    tb(s, Inches(8.3), Inches(1.4), Inches(4.45), Inches(0.4), "Recommended", 14, True, WHITE)
    tb(s, Inches(8.3), Inches(1.85), Inches(4.45), Inches(1.7), "BPSO لتخطيط الشبكة الإجمالي\nتغطية + كلفة + طاقة + تقارب + زمن", 16, True, WHITE)
    card(s, Inches(8.15), Inches(4.1), Inches(4.75), Inches(2.15), RGBColor(0xF8, 0xEE, 0xF0), AGA, 2)
    tb(s, Inches(8.3), Inches(4.3), Inches(4.45), Inches(1.7), "AGA تبقى جذابة عندما تكون الأولوية للاستقرار وقابلية التكرار.", 16, True, AGA)
    notes(s, "القرار", "BPSO للأداء الكلي. AGA للثبات.", "توصية مشروطة.", "لا تُلغِ AGA.", "أيهما للمشغّل؟", "BPSO للتوازن الكلي.")

    # 16 Takeaway
    s = blank(prs)
    header(s, prs, "Takeaway", "ثلاث رسائل من الفصل الرابع", 16)
    msgs = [
        ("01", "Real Data", "79,268 real network sites."),
        ("02", "Smart Optimization", "AGA + BPSO + Adaptive Constraint Repair."),
        ("03", "Actionable Decision", "BPSO: strongest overall trade-off under limited resources."),
    ]
    for i, (n, t, d) in enumerate(msgs):
        card(s, Inches(0.4 + i * 4.25), Inches(1.25), Inches(4.05), Inches(2.7))
        tb(s, Inches(0.55 + i * 4.25), Inches(1.4), Inches(3.75), Inches(0.5), n, 22, True, BPSO)
        tb(s, Inches(0.55 + i * 4.25), Inches(1.95), Inches(3.75), Inches(0.5), t, 16, True, INK)
        tb(s, Inches(0.55 + i * 4.25), Inches(2.55), Inches(3.75), Inches(1.1), d, 15, False, MUTED)
    card(s, Inches(0.4), Inches(4.2), Inches(12.5), Inches(2.15), INK)
    tb(
        s,
        Inches(0.7),
        Inches(4.55),
        Inches(11.9),
        Inches(1.5),
        "الهدف ليس نشر مزيد من المواقع — بل نشر المواقع الصحيحة.\nThe objective is not to deploy more sites — it is to deploy the right sites.",
        20,
        True,
        WHITE,
        "center",
    )
    notes(s, "الختام", "ثلاث رسائل ثم الجملة الأخيرة فقط.", "المواقع الصحيحة لا المزيد.", "لا تفتح الفصل الخامس.", "أهم مساهمة؟", "إطار بيانات + نموذج + خوارزميتان + إصلاح تكيفي → قرار هندسي.")

    prs.save(PPTX_PATH)
    print("PPTX:", PPTX_PATH)


NOTES = r"""الفصل الرابع — سكربت الإلقاء وأسئلة اللجنة
المصدر الوحيد: الفصل الرابع من الأطروحة. لا أرقام إضافية.

============================================================
أرقام يجب حفظها
============================================================
79,268 مواقع | 21,356 2G | 27,902 3G | 30,010 4G مرشحة
مشغّلان | 62% حضري | 38% ريفي
P/S = 150 | T = 300 | 15 تشغيلاً
S1=5,000 | S2=10,000 | S3=15,000
BPSO: 95.14% | 130.5 M$ | 76.4 MWh | 118 s | ±0.42 / ±2.9 / ±2.1
AGA:  94.90% | 139.6 M$ | 80.9 MWh | 142 s | ±0.36 / ±2.5 / ±1.8
Std.GA: 94.75% | 143.8 M$ | 84.1 MWh | 118 s
Random: 88.35% | 150.3 M$ | 91.5 MWh | 25 s
فروق BPSO مقابل AGA: +0.24% تغطية | ~6.5% كلفة | ~5.5% طاقة
ميزانية→تغطية AGA/BPSO: 80: 88.7/89.4 | 100: 91.9/92.8 | 120: 93.8/94.5 | 140: 94.8/95.1

============================================================
سكربت الشرائح (30–60 ثانية)
============================================================
01 التحدي
أقول: شبكة حقيقية 79,268 موقعاً، موارد محدودة، ثلاثة أهداف. السؤال: أي المواقع نرقّي؟
لا أتوسع: لا نتائج. الخريطة تخطيطية.

02 المسار
أقول: بيانات → تنظيف → GIS → نموذج → AGA/BPSO → تجارب → قرار.

03 البيانات
أقول أرقام الجدول 3. 4G هي قاعدة NSA.

04 الصعوبة
أقول: قرار ثنائي، فضاء 2^30,010، NP-hard، قيود صلبة. الاستدلال الفوقي ضرورة.

05 النموذج
أقول: max F = w1 Coverage − w2 Cost − w3 Energy، والأوزان مجموعها 1 بعد Min-Max. قيود B وK وCmin.

06 خوارزميتان
أقول: AGA تطورية، BPSO سربية، إصلاح مشترك، حجم 150 وتكرار 300.

07 الإصلاح (الشريحة المحورية)
أقول: غير الصالح يُصلَح. تجاوز B → حذف أدنى تغطية/تكلفة. نقص Cmin → إضافة الأعلى. بعد كل تحديث.

08 التجارب
أقول: 150 / 300 / 15. سيناريوهات 5k و10k و15k. النتائج التفصيلية لـ S2.

09 النتيجة
أقول أرقام BPSO ثم AGA ثم الفروق الثلاثة. لا أعلن تفوقاً مطلقاً.

10 الاستقرار
أقول: BPSO أفضل متوسطاً، AGA أقل تشتتاً.

11 التقارب
أقول: رسم توضيحي للشكل 14. أسرع 80، استقرار ~200، 300 كافٍ.

12 باريتو
أقول: توضيح مفهومي للشكل 15. BPSO أقرب، AGA أكثر تنوعاً.

13 الميزانية
أقول أرقام الجدول 12 ثم العائد المتناقص بعد ≈140 M$.

14 خطوط الأساس
أقول Random وStd.GA مقابل المقترحتين. الأزمنة 118 / 142 / 25.

15 القرار
أقول: BPSO للتخطيط الإجمالي. AGA إن طُلب الثبات.

16 الختام
أقول ثلاث رسائل ثم: الهدف ليس نشر مزيد من المواقع بل نشر المواقع الصحيحة.

============================================================
أسئلة اللجنة وإجابات قصيرة من الفصل فقط
============================================================
Q1 لماذا BPSO وAGA؟
لأن المسألة ثنائية واسعة النطاق ومتعددة الأهداف. الأدبيات تعتمد GA وPSO. طبّقناهما على البيانات نفسها مع إصلاح مشترك لمقارنة منهجية.

Q2 لماذا لا MILP أو حل دقيق؟
الفصل يصنّف المسألة NP-hard وغير خطية بمتغيرات ثنائية وعلى نطاق واسع. الحل الدقيق غير عملي؛ لذلك metaheuristics ضرورة.

Q3 ما أهمية Constraint Repair؟
تُطبَّق بعد كل تحديث. تجاوز الميزانية: إزالة أدنى مساهمة تغطية/تكلفة. نقص التغطية: إضافة الأعلى حتى Cmin. الحلول تُصلَح لا تُهدر، والآلية مشتركة.

Q4 لماذا BPSO أفضل في النتائج النهائية؟
الجدول 11 (S2): تغطية أعلى بـ 0.24%، كلفة أقل بنحو 6.5%، طاقة أفضل بنحو 5.5%، وتقارب أسرع و118 ثانية مقابل 142.

Q5 لماذا AGA أكثر استقراراً؟
الجدولان 9 و10: انحراف معياري أدنى في التغطية والكلفة والطاقة. الفصل يربط ذلك بتطور متدرج وتكيف الطفرة مقابل تقارب السرب الأسرع.

Q6 لماذا 300 تكرار؟
تحليل الشكل 14: الاستقرار بعد نحو 200 تكرار؛ إذن T=300 كافٍ عملياً.

Q7 لماذا 15 تشغيلاً؟
لتقدير المتوسط والانحراف عبر بذور عشوائية مختلفة. النتائج المبلَّغة متوسطات.

Q8 لماذا 5,000 / 10,000 / 15,000؟
ثلاثة سيناريوهات لحجم/تمويل منخفض ومتوسط ومرتفع (الجدول 6).

Q9 ما معنى Pareto Front هنا؟
جبهة المفاضلة بين التغطية والكلفة. BPSO أقرب إلى المنطقة المثلى؛ AGA أكثر تنوعاً. كلاهما يتفوق على الأساس.

Q10 هل تُعمَّم النتائج على شبكات أخرى؟
الأرقام خاصة بالشبكة السورية المدروسة. المنهج (بيانات تشغيلية + GIS + نموذج مقيد + إصلاح تكيفي) قابل للنقل، لكن الأرقام لا تُنقل دون إعادة معايرة. الفصل لا يدّعي تعميماً آلياً.

Q11 ما أهمية 4G في NSA؟
4G/LTE هي البنية الأكثر جاهزية لدعم 5G وفق المعمارية غير المستقلة ومعايير 3GPP. لذلك N = 30,010 مرشحاً.

Q12 ما أهم مساهمة علمية في الفصل؟
إطار متكامل: بيانات تشغيلية واسعة، نموذج ثلاثي الأهداف مقيد، خوارزميتان مع إصلاح قيود تكيفي، ومقارنة تجريبية تفضي إلى قرار هندسي تحت موارد محدودة.
"""


if __name__ == "__main__":
    (ROOT / "deliverables").mkdir(parents=True, exist_ok=True)
    NOTES_PATH.write_text(NOTES, encoding="utf-8")
    build()
    print("Notes:", NOTES_PATH)
    print("Charts:", CHARTS)
